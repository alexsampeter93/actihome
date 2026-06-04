from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
    KeepTogether,
)
from reportlab.pdfbase.pdfmetrics import stringWidth


DOCX_PATH = Path("docs/ActiHome_Guia_Tecnica_Estudio.docx")
PDF_PATH = Path("docs/ActiHome_Guia_Tecnica_Estudio.pdf")


def iter_block_items(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            for paragraph in document.paragraphs:
                if paragraph._p is child:
                    yield ("paragraph", paragraph)
                    break
        elif child.tag == qn("w:tbl"):
            for table in document.tables:
                if table._tbl is child:
                    yield ("table", table)
                    break


def escape(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def paragraph_text(paragraph):
    text = paragraph.text.strip()
    return text


def is_code_table(table):
    if len(table.rows) != 1 or len(table.columns) != 1:
        return False
    text = table.cell(0, 0).text
    code_markers = ["public ", "@Entity", "CREATE TABLE", "if (", "Usuario pulsa", "SpringApplication"]
    return "\n" in text or any(marker in text for marker in code_markers)


def split_code_lines(code, max_chars=88):
    lines = []
    for line in code.splitlines():
        if len(line) <= max_chars:
            lines.append(line)
            continue
        current = line
        while len(current) > max_chars:
            cut = current.rfind(" ", 0, max_chars)
            if cut < 30:
                cut = max_chars
            lines.append(current[:cut])
            current = "    " + current[cut:].lstrip()
        lines.append(current)
    return "\n".join(lines)


def make_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="GuideTitle",
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=26,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#1F4E79"),
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideSubtitle",
            fontName="Helvetica",
            fontSize=11,
            leading=15,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#666666"),
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideH1",
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=19,
            textColor=colors.HexColor("#1F4E79"),
            spaceBefore=12,
            spaceAfter=7,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideH2",
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#2F5597"),
            spaceBefore=8,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideBody",
            fontName="Helvetica",
            fontSize=9.6,
            leading=13,
            alignment=TA_LEFT,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideCode",
            fontName="Courier",
            fontSize=7.5,
            leading=9.4,
            textColor=colors.HexColor("#222222"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="GuideSmall",
            fontName="Helvetica",
            fontSize=8.4,
            leading=11,
        )
    )
    return styles


def draw_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawString(1.7 * cm, 1.05 * cm, "ActiHome - Guia tecnica de estudio")
    canvas.drawRightString(A4[0] - 1.7 * cm, 1.05 * cm, f"Pagina {doc.page}")
    canvas.restoreState()


def add_docx_table(story, table, styles):
    rows = []
    for row in table.rows:
        pdf_row = []
        for cell in row.cells:
            text = escape(cell.text.strip())
            pdf_row.append(Paragraph(text, styles["GuideSmall"]))
        rows.append(pdf_row)

    if not rows:
        return

    col_count = len(rows[0])
    available_width = A4[0] - 3.4 * cm
    col_widths = [available_width / col_count] * col_count
    pdf_table = Table(rows, colWidths=col_widths, hAlign="LEFT", repeatRows=1)
    pdf_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEBF7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F1F1F")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#BFBFBF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(pdf_table)
    story.append(Spacer(1, 0.18 * cm))


def add_code_table(story, table, styles):
    code = split_code_lines(table.cell(0, 0).text.strip())
    code_para = Paragraph(escape(code), styles["GuideCode"])
    pdf_table = Table([[code_para]], colWidths=[A4[0] - 3.4 * cm], hAlign="LEFT")
    pdf_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F6F8")),
                ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D3DD")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(KeepTogether([pdf_table, Spacer(1, 0.16 * cm)]))


def build_pdf():
    source = Document(DOCX_PATH)
    styles = make_styles()
    story = []

    for kind, block in iter_block_items(source):
        if kind == "paragraph":
            text = paragraph_text(block)
            if not text:
                story.append(Spacer(1, 0.08 * cm))
                continue
            style_name = block.style.name
            if style_name == "Title":
                story.append(Paragraph(escape(text), styles["GuideTitle"]))
            elif style_name == "Heading 1":
                story.append(Paragraph(escape(text), styles["GuideH1"]))
            elif style_name == "Heading 2":
                story.append(Paragraph(escape(text), styles["GuideH2"]))
            elif style_name == "List Number":
                story.append(Paragraph("&#8226; " + escape(text), styles["GuideBody"]))
            elif style_name == "List Bullet":
                story.append(Paragraph("&#8226; " + escape(text), styles["GuideBody"]))
            else:
                if "Como se construye una aplicacion" in text:
                    story.append(Paragraph(escape(text), styles["GuideSubtitle"]))
                else:
                    story.append(Paragraph(escape(text), styles["GuideBody"]))
        elif kind == "table":
            if is_code_table(block):
                add_code_table(story, block, styles)
            else:
                add_docx_table(story, block, styles)

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        rightMargin=1.7 * cm,
        leftMargin=1.7 * cm,
        topMargin=1.55 * cm,
        bottomMargin=1.7 * cm,
        title="ActiHome - Guia tecnica de estudio",
        author="Codex",
    )
    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


if __name__ == "__main__":
    build_pdf()
